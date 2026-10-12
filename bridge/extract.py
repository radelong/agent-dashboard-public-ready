"""Bounded, non-executing extraction. Originals are preserved in hosted storage."""
import sys, json, base64, io, zipfile, re, csv
from html.parser import HTMLParser

MAX_TEXT = 1_500_000
sections = []
warnings = []
def add(text, location):
    text = re.sub(r'[ \t]+', ' ', str(text)).strip()
    if text:
        sections.append({'text': text, 'location': location})
        if sum(len(s['text']) for s in sections) > MAX_TEXT:
            raise ValueError('Extracted text exceeds 1.5 million characters; split this document.')

class TextHTML(HTMLParser):
    def __init__(self):
        super().__init__(); self.parts = []; self.skip = 0
    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'style'): self.skip += 1
        if tag in ('p', 'div', 'br', 'li', 'h1', 'h2', 'h3', 'tr'): self.parts.append('\n')
    def handle_endtag(self, tag):
        if tag in ('script', 'style'): self.skip = max(0, self.skip - 1)
    def handle_data(self, data):
        if not self.skip: self.parts.append(data)

def main():
    data = json.load(sys.stdin); raw = base64.b64decode(data['base64'], validate=True)
    ext = data['extension']; source = io.BytesIO(raw)
    if ext in ('docx', 'pptx', 'xlsx'):
        with zipfile.ZipFile(source) as archive:
            if len(archive.infolist()) > 10000 or sum(i.file_size for i in archive.infolist()) > 100_000_000:
                raise ValueError('Office archive expands beyond safe extraction limits.')
        source.seek(0)
    if ext == 'pdf':
        from pypdf import PdfReader
        reader = PdfReader(source)
        if reader.is_encrypted: raise ValueError('Password-protected PDF: submit an unlocked copy.')
        if len(reader.pages) > 1000: raise ValueError('PDF has more than 1,000 pages; split it.')
        blank = []
        for n, page in enumerate(reader.pages, 1):
            if len(page.get_contents().get_data()) > 20_000_000 if page.get_contents() else False:
                raise ValueError('PDF page expands beyond extraction limits.')
            text = page.extract_text() or ''
            if text.strip(): add(text, f'Page {n}')
            else: blank.append(n)
        if blank: warnings.append('Pages with no selectable text (OCR needed if scanned): ' + ', '.join(map(str, blank[:30])))
    elif ext == 'docx':
        from docx import Document
        doc = Document(source)
        # Preserve paragraph/table reading order; page numbers cannot be inferred from DOCX XML.
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        for n, block in enumerate(doc.element.body, 1):
            if block.tag.endswith('}p'): add(Paragraph(block, doc).text, f'Block {n}')
            elif block.tag.endswith('}tbl'):
                table = Table(block, doc); add('\n'.join(' | '.join(c.text for c in row.cells) for row in table.rows), f'Table at block {n}')
    elif ext == 'pptx':
        from pptx import Presentation
        for n, slide in enumerate(Presentation(source).slides, 1):
            parts = []
            for shape in slide.shapes:
                if shape.has_text_frame: parts.append(shape.text)
                if shape.has_table: parts.extend(' | '.join(c.text for c in row.cells) for row in shape.table.rows)
            add('\n'.join(parts), f'Slide {n}')
    elif ext == 'xlsx':
        from openpyxl import load_workbook
        book = load_workbook(source, read_only=True, data_only=True)
        warnings.append('Spreadsheet uses stored cell values; formulas are not recalculated. Images and charts are not extracted.')
        count = 0
        for sheet in book:
            for n, row in enumerate(sheet.iter_rows(values_only=True), 1):
                count += len(row)
                if count > 500000: raise ValueError('Spreadsheet exceeds 500,000 cells; submit selected sheets.')
                add(' | '.join('' if v is None else str(v) for v in row), f'{sheet.title}, row {n}')
        book.close()
    else:
        try: text = raw.decode('utf-8-sig')
        except UnicodeDecodeError: raise ValueError('Text files must use UTF-8 encoding.')
        if '\x00' in text: raise ValueError('This file is not a supported text document.')
        if ext in ('html', 'htm'):
            parser = TextHTML(); parser.feed(text); text = ''.join(parser.parts)
        lines = text.splitlines()
        for start in range(0, len(lines), 40): add('\n'.join(lines[start:start+40]), f'Lines {start+1}–{min(start+40,len(lines))}')
    if not sections: raise ValueError('No readable text found. Scanned documents need OCR before submission.')
    print(json.dumps({'sections': sections, 'warnings': warnings}, ensure_ascii=False))

try: main()
except Exception as e:
    print(json.dumps({'error': str(e)[:1000]})); sys.exit(1)
